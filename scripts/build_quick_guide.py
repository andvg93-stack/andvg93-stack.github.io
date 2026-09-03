from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_activity_document import (
    BORDER,
    GREEN,
    GREEN_LIGHT,
    GOLD_LIGHT,
    INK,
    WHITE,
    add_numbered,
    create_numbering_instance,
    install_numbering,
    set_cell_borders,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
    set_table_geometry,
)


OUTPUT = Path(__file__).resolve().parents[1] / "deliverables" / "Transformando_Territorios_Guia_Operativa_4_Sep_2026.docx"


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("4 sep 2026  |  Página ")
    set_run_font(run, size=8, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend([begin, instruction, separate, value, end])
    set_run_font(field_run, size=8, color="666666")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_label(document, text, before=6, after=3):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK, bold=True)
    return paragraph


def add_schedule(document):
    headers = ["Hora", "Momento", "Qué hacer", "No olvidar"]
    rows = [
        ("10:00-10:10", "Apertura", "Bienvenida, propósito, ruta y dos acuerdos de participación.", "Nombrar la Cátedra Social Solidaria y el tema central."),
        ("10:10-10:30", "Clima y resiliencia", "Aporte experto sobre adaptación local y resiliencia comunitaria.", "Pedir un ejemplo relacionado con San Adolfo o Acevedo."),
        ("10:30-10:50", "Sistemas de Información Geográfica", "Explicar cómo los SIG apoyan monitoreo y prospección climática.", "Conectar mapas, datos y decisiones territoriales."),
        ("10:50-11:05", "Jóvenes de Ambiente - CAM", "Presentar la estrategia y opciones de participación juvenil.", "Dejar clara una ruta de vinculación o contacto."),
        ("11:05-11:15", "Preguntas", "Recoger dos o tres preguntas juveniles y moderar respuestas breves.", "Relacionar las respuestas con el territorio."),
        ("11:15-11:25", "Pausa y montaje", "Organizar grupos y abrir el simulador en 2026.", "Probar imagen, internet y respaldo local."),
        ("11:25-11:40", "Café 2035", "Explicar línea temporal, leyenda, indicadores y expansión cafetera estimada.", "No es una predicción de predios ni una autorización."),
        ("11:40-12:10", "Exploración", "En grupos, comparar 2026, 2030 y 2035 y anotar dos cambios.", "Mirar mapa e indicadores; pausar en cada fecha."),
        ("12:10-12:30", "Impactos", "Elegir agua, suelo, biodiversidad, clima o emisiones y priorizar un efecto.", "Distinguir dato, experiencia local y pregunta por verificar."),
        ("12:30-12:50", "Acción juvenil", "Definir problema, acción, aliado, primer paso e indicador sencillo.", "La acción debe poder iniciar en tres meses."),
        ("12:50-1:00", "Cierre", "Un minuto por grupo; acordar enlace, próximo paso y evaluación rápida.", "Cerrar a tiempo y recoger todas las fichas."),
    ]
    table = document.add_table(rows=1, cols=4)
    widths = [1250, 2140, 3320, 2650]
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, GREEN_LIGHT)
        set_cell_borders(cell, BORDER, "6")
        set_cell_margins(cell, top=70, bottom=70)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=9, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            set_cell_borders(cell, BORDER, "5")
            set_cell_margins(cell, top=58, bottom=58)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if idx == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            set_run_font(run, size=8.35, color="252525", bold=(idx == 1))
    set_table_geometry(table, widths)


def add_checklist(document, items):
    table = document.add_table(rows=0, cols=2)
    widths = [420, 8940]
    for item in items:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_borders(cell, WHITE, "0")
            set_cell_margins(cell, top=22, bottom=22, start=60, end=80)
        mark = cells[0].paragraphs[0]
        mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mark.paragraph_format.space_after = Pt(0)
        square = mark.add_run("☐")
        set_run_font(square, name="Arial", size=11, color=GREEN)
        text = cells[1].paragraphs[0]
        text.paragraph_format.space_after = Pt(0)
        run = text.add_run(item)
        set_run_font(run, size=9, color="252525")
    set_table_geometry(table, widths)


def add_note(document, label, text, fill=GOLD_LIGHT):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, GREEN, "6")
    set_cell_margins(cell, top=70, bottom=70)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    first = paragraph.add_run(label + "  ")
    set_run_font(first, size=9, color=INK, bold=True)
    rest = paragraph.add_run(text)
    set_run_font(rest, size=9, color="252525")


def build_quick_guide():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.footer_distance = Inches(0.38)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(37, 37, 37)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    _, decimal_abstract_id = install_numbering(document)

    line = document.add_paragraph()
    line.paragraph_format.space_after = Pt(3)
    title = line.add_run("Transformando Territorios")
    set_run_font(title, size=15, color=INK, bold=True)
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(7)
    meta_run = meta.add_run("Viernes 4 de septiembre de 2026  |  10:00 a. m.  |  San Adolfo, Acevedo - Huila")
    set_run_font(meta_run, size=9.5, color="555555", bold=True)
    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(7)
    intro_run = intro.add_run("Ruta propuesta de 10:00 a. m. a 1:00 p. m. Confirmar antes de iniciar el orden y tema de cada persona invitada.")
    set_run_font(intro_run, size=9.5, color="252525")

    add_schedule(document)

    document.add_page_break()
    add_label(document, "Antes de empezar", before=0)
    add_checklist(
        document,
        [
            "Confirmar hora de cierre, orden de intervenciones y quién controla el tiempo.",
            "Tener listo computador, cargador, proyector, sonido, extensión y adaptadores.",
            "Abrir el simulador y comprobar Play, pausa, línea temporal, cauces, municipios e indicadores.",
            "Tener a mano el respaldo local y capturas impresas de 2026, 2030 y 2035.",
            "Separar fichas, marcadores y materiales por grupo; definir relatoría y vocerías.",
            "Acordar quién toma fotografías y verificar las autorizaciones correspondientes.",
        ],
    )

    add_label(document, "Durante el simulador")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for step in [
        "Detenerse en 2026: pedir dos rasgos del escenario inicial.",
        "Avanzar a 2030: preguntar qué cambió en el mapa y en las barras de salud territorial.",
        "Llegar a 2035: identificar el efecto más importante y a quién puede afectar.",
        "Contrastar: separar estimación del simulador, experiencia local y pregunta que requiere verificación.",
        "Convertir el hallazgo en acción: problema + acción + aliado + primer paso + indicador.",
    ]:
        paragraph = add_numbered(document, step, decimal_id)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=9, color="252525")

    add_label(document, "Frases que no se deben olvidar")
    add_note(document, "AL PRESENTAR EL SIMULADOR", "Este es un escenario didáctico estimado. No predice el futuro de un predio ni autoriza expansión agrícola.", fill=GREEN_LIGHT)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    add_note(document, "AL PASAR A LA ACCIÓN", "No buscamos resolver todo hoy; buscamos un primer paso que pueda comenzar en los próximos tres meses.")

    add_label(document, "Si el tiempo se atrasa")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run("Proteger como mínimo 60 minutos finales: 10 min de explicación, 20 min de exploración, 15 min de impactos, 10 min de acción y 5 min de cierre. Reducir preguntas plenarias; no eliminar el cierre.")
    set_run_font(run, size=9, color="252525")

    add_label(document, "Contingencias rápidas")
    add_checklist(
        document,
        [
            "Sin internet: usar el respaldo local; el mapa y los indicadores esenciales deben seguir disponibles.",
            "Sin proyector o energía: trabajar con capturas impresas de 2026, 2030 y 2035.",
            "Pocas voces dominan: hacer una ronda de 30 segundos y usar vocerías rotativas.",
            "Surge un dato dudoso: registrarlo como pregunta por verificar; no improvisar una respuesta.",
        ],
    )

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core = document.core_properties
    core.title = "Transformando Territorios - Guía operativa"
    core.subject = "Cronograma de facilitación para la jornada del 4 de septiembre de 2026"
    core.author = "Equipo de trabajo - Transformando Territorios"
    core.keywords = "San Adolfo, Acevedo, juventud, resiliencia climática, Café 2035"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_quick_guide()
