from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "deliverables" / "Transformando_Territorios_Guia_Metodologica.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203B3A"
GREEN = "2F6B57"
GREEN_LIGHT = "EAF2ED"
GOLD = "C9912B"
GOLD_LIGHT = "FBF3DE"
GRAY = "5D6866"
LIGHT_GRAY = "F2F4F7"
BORDER = "CBD5D2"
WHITE = "FFFFFF"
RED = "9B1C1C"


def set_run_font(run, name="Calibri", size=11, color="222222", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=GRAY)
    for field_name in ("PAGE",):
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = field_name
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Calibri")
    run_fonts.set(qn("w:hAnsi"), "Calibri")
    props.append(run_fonts)
    tint = OxmlElement("w:color")
    tint.set(qn("w:val"), color)
    props.append(tint)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def install_numbering(document):
    numbering = document.part.numbering_part.element
    existing_abstract = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    existing_nums = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    base_abstract = max(existing_abstract or [0]) + 1
    base_num = max(existing_nums or [0]) + 1

    def create(abstract_id, num_id, fmt, text, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{0xA1352000 + abstract_id:08X}")
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        template = OxmlElement("w:tmpl")
        template.set(qn("w:val"), f"{0xC0FFEE00 + abstract_id:08X}")
        abstract.append(template)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        level.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        level.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        level.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            level.append(r_pr)
        abstract.append(level)
        # OOXML requires every abstract numbering definition to precede the
        # concrete <w:num> instances. Appending an abstract after the template's
        # existing nums makes Word repair the whole list table on open.
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(numbering.index(first_num), abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    create(base_abstract, base_num, "decimal", "%1.")
    # The default DOCX template already contains a native List Bullet definition
    # (numId 1). Reusing it maximizes compatibility across Word renderers.
    return 1, base_abstract


def create_numbering_instance(document, abstract_id):
    numbering = document.part.numbering_part.element
    existing_nums = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_nums or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def set_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    identity = OxmlElement("w:numId")
    identity.set(qn("w:val"), str(num_id))
    num_pr.extend([level, identity])


def add_bullet(document, text, num_id, bold_prefix=None):
    paragraph = document.add_paragraph()
    set_numbering(paragraph, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_numbered(document, text, num_id):
    paragraph = document.add_paragraph()
    set_numbering(paragraph, num_id)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(document, label, text, fill=GREEN_LIGHT, accent=GREEN):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, accent, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    label_run = p.add_run(label + "  ")
    set_run_font(label_run, size=10, color=accent, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document, headers, rows, widths, header_fill=LIGHT_GRAY, alignments=None):
    table = document.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if not alignments else alignments[index]
        run = p.add_run(header)
        set_run_font(run, size=9.5, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            set_cell_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if alignments:
                p.alignment = alignments[index]
            run = p.add_run(value)
            set_run_font(run, size=9.5, color="303938")
    set_table_geometry(table, widths)
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    set_keep_with_next(paragraph)
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_prompt_box(document, number, title, prompt, lines=3):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{number}. {title}")
    set_run_font(r, size=11, color=GREEN, bold=True)
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAF9")
    set_cell_borders(cell, BORDER, "6")
    prompt_p = cell.paragraphs[0]
    prompt_p.paragraph_format.space_after = Pt(5)
    prompt_r = prompt_p.add_run(prompt)
    set_run_font(prompt_r, size=9.5, color=GRAY, italic=True)
    for _ in range(lines):
        blank = cell.add_paragraph(" ")
        blank.paragraph_format.space_after = Pt(4)
        p_pr = blank._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), "D8E0DD")
        borders.append(bottom)
        p_pr.append(borders)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False
    document.settings.odd_and_even_pages_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("TRANSFORMANDO TERRITORIOS")
        set_run_font(r, size=8.5, color=GREEN, bold=True)
        r2 = p.add_run("  |  Guía metodológica")
        set_run_font(r2, size=8.5, color=GRAY)
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "D9E4DF")
        borders.append(bottom)
        p_pr.append(borders)

    for footer in (section.footer, section.even_page_footer):
        add_page_field(footer.paragraphs[0])

    core = document.core_properties
    core.title = "Transformando Territorios: Juventud, Innovación y Acción Climática desde el Huila"
    core.subject = "Guía metodológica para el equipo de trabajo"
    core.author = "Equipo de trabajo - Transformando Territorios"
    core.keywords = "Huila, juventud, acción climática, territorio, café 2035"


def build_document():
    document = Document()
    configure_document(document)
    bullet_id, decimal_abstract_id = install_numbering(document)

    # Portada / workshop agenda header pattern.
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("GUÍA METODOLÓGICA")
    set_run_font(r, size=10, color=GOLD, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    title = p.add_run("Transformando Territorios:")
    set_run_font(title, size=27, color=INK, bold=True)
    title.add_break()
    title2 = p.add_run("Juventud, Innovación y Acción Climática desde el Huila")
    set_run_font(title2, size=22, color=GREEN, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Propuesta operativa para la planeación, facilitación y seguimiento de una experiencia educativa con jóvenes del sur del Huila.")
    set_run_font(r, size=12.5, color=GRAY, italic=True)

    metrics = document.add_table(rows=1, cols=3)
    set_table_geometry(metrics, [3120, 3120, 3120])
    metric_values = [
        ("PÚBLICO", "Jóvenes de 18 a 25 años"),
        ("DURACIÓN PROPUESTA", "3 horas presenciales"),
        ("EJE PRÁCTICO", "Simulador Café 2035"),
    ]
    for index, (label, value) in enumerate(metric_values):
        cell = metrics.rows[0].cells[index]
        set_cell_shading(cell, GREEN_LIGHT if index != 1 else GOLD_LIGHT)
        set_cell_borders(cell, WHITE, "8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        label_run = p.add_run(label)
        set_run_font(label_run, size=8.5, color=GREEN if index != 1 else "8A651B", bold=True)
        value_p = cell.add_paragraph()
        value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_p.paragraph_format.space_after = Pt(0)
        value_run = value_p.add_run(value)
        set_run_font(value_run, size=10.5, color=INK, bold=True)

    document.add_paragraph().paragraph_format.space_after = Pt(8)
    add_callout(
        document,
        "PROPÓSITO",
        "Conectar la experiencia cotidiana de las juventudes con los cambios del clima y del territorio, interpretar colectivamente un escenario de evolución de la huella cafetera y convertir los hallazgos en propuestas juveniles de acción local.",
    )

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Documento de trabajo para revisión del equipo")
    set_run_font(r, size=10.5, color=INK, bold=True)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Versión 1.0 | Septiembre de 2026")
    set_run_font(r, size=10, color=GRAY)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Ámbito territorial: municipios del sur del departamento del Huila")
    set_run_font(r, size=10, color=GRAY)

    document.add_page_break()

    add_heading(document, "1. Resumen ejecutivo", 1)
    add_body(
        document,
        "La actividad propone una experiencia participativa en la que jóvenes de 18 a 25 años analizan cómo el cambio climático puede transformar el territorio cafetero del Huila entre 2026 y 2035. El simulador Café 2035 funciona como detonante de conversación: permite observar la evolución temporal de la huella cafetera estimada y su relación con emisiones de CO₂ equivalente, fuentes hídricas, salud del suelo, biodiversidad y resiliencia climática.",
    )
    add_body(
        document,
        "La sesión no busca enseñar una predicción exacta ni tomar decisiones productivas. Su propósito es fortalecer la lectura crítica del territorio, la comprensión de interdependencias socioambientales y la capacidad juvenil para formular acciones viables de adaptación, cuidado y comunicación climática.",
    )

    add_heading(document, "2. Ficha técnica", 1)
    add_table(
        document,
        ["Elemento", "Definición propuesta"],
        [
            ("Nombre", "Transformando Territorios: Juventud, Innovación y Acción Climática desde el Huila"),
            ("Población", "Jóvenes de 18 a 25 años; se recomienda un grupo de 20 a 35 participantes."),
            ("Cobertura", "Sur del Huila, con énfasis en Pitalito, Acevedo, Elías, Isnos, Oporapa, Palestina, Saladoblanco, San Agustín y Timaná."),
            ("Modalidad", "Presencial, apoyada por proyección del simulador web y trabajo colaborativo."),
            ("Duración", "180 minutos, incluida una pausa activa de 15 minutos."),
            ("Producto central", "Una propuesta juvenil de acción territorial por grupo, con problema priorizado, actores, primer paso e indicador sencillo."),
            ("Equipo mínimo", "Una persona coordinadora, una facilitadora principal, una persona de apoyo técnico y una persona de relatoría/logística."),
        ],
        [1900, 7460],
        header_fill=GREEN_LIGHT,
    )

    add_heading(document, "3. Justificación", 1)
    add_body(
        document,
        "El café hace parte de la identidad productiva, cultural y paisajística del Huila. Al mismo tiempo, la variación de temperatura y precipitación, la presión sobre el agua, el suelo y los ecosistemas, y los cambios en la aptitud climática plantean retos que no pueden abordarse únicamente desde información técnica. Es necesario crear espacios en los que las juventudes relacionen los datos con sus propias experiencias y reconozcan su capacidad para intervenir en el presente del territorio.",
    )
    add_body(
        document,
        "La metodología combina visualización, diálogo y diseño de acciones. El recorrido temporal del simulador facilita una lectura comparativa y didáctica; las preguntas de facilitación permiten contrastar la proyección con saberes locales; y el laboratorio final conduce a propuestas concretas que pueden alimentar procesos comunitarios, educativos o institucionales.",
    )

    add_heading(document, "4. Objetivos", 1)
    add_heading(document, "4.1 Objetivo general", 2)
    add_body(
        document,
        "Fortalecer en jóvenes del sur del Huila la comprensión crítica de las relaciones entre cambio climático, expansión de la huella cafetera y salud del territorio, promoviendo propuestas juveniles de acción climática local.",
    )
    add_heading(document, "4.2 Objetivos específicos", 2)
    for item in [
        "Reconocer percepciones y cambios observados por las juventudes en su territorio.",
        "Interpretar un escenario didáctico 2026-2035 y sus posibles efectos sobre agua, suelos, biodiversidad, clima y emisiones.",
        "Diferenciar entre proyección estimada, evidencia observada y conocimiento local.",
        "Identificar riesgos, oportunidades y actores relevantes en los municipios participantes.",
        "Formular acciones juveniles realizables, medibles y conectadas con el contexto territorial.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "5. Resultados esperados", 1)
    for item in [
        "Un mapa colectivo de percepciones sobre cambios climáticos y territoriales.",
        "Una ficha de lectura del simulador por grupo, con hallazgos para 2026, 2030 y 2035.",
        "Entre cuatro y seis propuestas juveniles de acción territorial.",
        "Un registro de compromisos, actores aliados y próximos pasos.",
        "Una evaluación rápida de aprendizaje y pertinencia de la actividad.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "6. Principios metodológicos", 1)
    principles = [
        ("Participación activa.", "La juventud interpreta, pregunta, propone y toma la palabra; no se limita a recibir información."),
        ("Lectura situada.", "Los datos se contrastan con experiencias, memorias y conocimientos de los municipios."),
        ("Rigor con lenguaje claro.", "Se explican supuestos, alcances y limitaciones sin perder la dimensión pedagógica."),
        ("Enfoque de soluciones.", "La discusión sobre impactos termina en alternativas concretas y realizables."),
        ("Cuidado e inclusión.", "Se distribuyen turnos de participación y se evita responsabilizar individualmente a las comunidades por problemas estructurales."),
    ]
    for label, detail in principles:
        add_bullet(document, label + " " + detail, bullet_id, bold_prefix=label)

    add_heading(document, "7. Preparación previa", 1)
    add_heading(document, "7.1 Decisiones que debe cerrar el equipo", 2)
    for item in [
        "Fecha, municipio, lugar, capacidad y condiciones de accesibilidad.",
        "Número estimado de participantes y mecanismo de convocatoria.",
        "Disponibilidad de internet, proyector, sonido, energía y plan de respaldo.",
        "Composición de los grupos y criterio para asegurar diversidad territorial y de género.",
        "Persona responsable de fotografías o testimonios y consentimiento correspondiente.",
        "Ruta de seguimiento para las propuestas producidas durante la sesión.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "7.2 Alistamiento del simulador", 2)
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Abrir y probar el simulador en el computador que se utilizará en la actividad.",
        "Verificar la reproducción automática, la pausa, la selección manual del tiempo y la visualización de municipios, cauces e indicadores.",
        "Ensayar la comparación entre enero de 2026, un punto intermedio cercano a 2030 y diciembre de 2035.",
        "Confirmar el funcionamiento del respaldo local sin teselas externas si la conectividad es inestable.",
        "Recordar al equipo que la huella cafetera y los indicadores son estimaciones didácticas; no corresponden a predicciones parcelarias ni autorizaciones de expansión.",
    ]:
        add_numbered(document, item, decimal_id)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    label = p.add_run("Acceso al simulador: ")
    set_run_font(label, bold=True, color=INK)
    add_hyperlink(p, "cafe-2035-huila.donchibcha.chatgpt.site", "https://cafe-2035-huila.donchibcha.chatgpt.site/")

    add_heading(document, "8. Agenda general de la sesión", 1)
    add_table(
        document,
        ["Tiempo", "Momento", "Propósito", "Producto o evidencia"],
        [
            ("0-15 min", "Bienvenida y acuerdos", "Presentar el propósito, crear condiciones de participación y activar conocimientos previos.", "Acuerdos visibles y expectativa inicial."),
            ("15-35 min", "El territorio que cambia", "Reconocer señales de cambio climático y territorial desde la experiencia juvenil.", "Mapa de percepciones y relatos breves."),
            ("35-50 min", "Claves para leer Café 2035", "Explicar el escenario, la línea temporal, la leyenda y los indicadores.", "Comprensión básica del uso y sus límites."),
            ("50-80 min", "Exploración guiada", "Comparar 2026, 2030 y 2035 por grupos.", "Ficha de hallazgos y preguntas."),
            ("80-95 min", "Pausa activa", "Descanso y conversación informal.", "Reagrupación."),
            ("95-125 min", "Lectura de impactos", "Relacionar cambios de la huella cafetera con agua, suelo, biodiversidad, clima y vida comunitaria.", "Prioridad territorial por grupo."),
            ("125-155 min", "Laboratorio de acción", "Diseñar una respuesta juvenil viable.", "Prototipo de acción territorial."),
            ("155-175 min", "Galería y compromisos", "Compartir, retroalimentar y elegir próximos pasos.", "Compromisos y posibles alianzas."),
            ("175-180 min", "Cierre y evaluación", "Recoger aprendizajes y valoración rápida.", "Evaluación de salida."),
        ],
        [1050, 2050, 3200, 3060],
        header_fill=GREEN_LIGHT,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(document, "9. Desarrollo metodológico", 1)
    add_heading(document, "Momento 1. Bienvenida y acuerdos | 15 minutos", 2)
    add_body(document, "Propósito: establecer un ambiente participativo, seguro y orientado a la acción.", bold_prefix="Propósito:")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Recibir a las y los participantes y explicar brevemente por qué se convoca el espacio.",
        "Presentar el objetivo y aclarar que no se evaluarán conocimientos técnicos previos.",
        "Construir tres acuerdos: escuchar sin descalificar, relacionar las ideas con el territorio y cuidar el tiempo de intervención.",
        "Aplicar una pregunta de entrada: ¿qué palabra describe hoy la relación entre juventud y territorio en el sur del Huila?",
    ]:
        add_numbered(document, item, decimal_id)
    add_callout(document, "CLAVE DE FACILITACIÓN", "Registrar las palabras iniciales para retomarlas en el cierre y mostrar cómo cambió la conversación.", fill=GOLD_LIGHT, accent="8A651B")

    add_heading(document, "Momento 2. El territorio que cambia | 20 minutos", 2)
    add_body(document, "Propósito: conectar la discusión climática con señales observadas en la vida cotidiana.", bold_prefix="Propósito:")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Ubicar un mapa impreso o dibujado del sur del Huila en una pared o mesa.",
        "Entregar notas adhesivas de dos colores: uno para cambios preocupantes y otro para respuestas u oportunidades existentes.",
        "Pedir que cada participante ubique una observación relacionada con lluvias, calor, agua, cultivos, bosques, suelos o movilidad.",
        "Agrupar observaciones repetidas y preguntar qué efectos tienen sobre jóvenes, familias y comunidades.",
    ]:
        add_numbered(document, item, decimal_id)
    add_body(document, "Preguntas orientadoras: ¿qué cambios se perciben desde hace algunos años?, ¿quiénes resultan más afectados?, ¿qué prácticas locales ya ayudan a responder?")

    add_heading(document, "Momento 3. Claves para leer Café 2035 | 15 minutos", 2)
    add_body(document, "Propósito: dar a todas las personas una base común para interpretar el simulador.", bold_prefix="Propósito:")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Presentar la línea temporal como la única variable controlable: se puede reproducir, pausar y elegir cualquier mes entre 2026 y 2035.",
        "Explicar la leyenda: café inicial o persistente, nueva expansión, retiro o pérdida de aptitud, áreas naturales/protegidas y cauces.",
        "Presentar los cinco indicadores de salud territorial y aclarar que expresan tendencias didácticas, no mediciones en tiempo real.",
        "Realizar una demostración corta sin interpretar todavía todos los resultados; la lectura corresponde a los grupos.",
    ]:
        add_numbered(document, item, decimal_id)
    add_callout(document, "MENSAJE OBLIGATORIO", "El simulador muestra un escenario estimado para aprender y conversar. No predice el futuro de un predio específico ni sustituye estudios técnicos o decisiones de ordenamiento.", fill="FCEDEC", accent=RED)

    add_heading(document, "Momento 4. Exploración guiada | 30 minutos", 2)
    add_body(document, "Propósito: identificar patrones, contrastes y preguntas a partir de tres cortes temporales.", bold_prefix="Propósito:")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Organizar grupos de cuatro a seis personas y asignar una relatoría, una vocería y una persona encargada de manejar la línea temporal.",
        "Observar el escenario inicial de 2026 y registrar dos características del territorio.",
        "Avanzar a un punto cercano a 2030 y registrar cambios visibles en la huella y los indicadores.",
        "Llegar a 2035, pausar y seleccionar el cambio que el grupo considera más importante.",
        "Contrastar lo observado con el mapa de percepciones construido al inicio.",
    ]:
        add_numbered(document, item, decimal_id)
    add_body(document, "Preguntas orientadoras: ¿dónde aparecen continuidades, expansiones o pérdidas?, ¿qué indicador cambia más?, ¿qué relación puede existir entre varios indicadores?, ¿qué no permite saber el simulador?")

    add_heading(document, "Momento 5. Lectura de impactos | 30 minutos", 2)
    add_body(document, "Cada grupo selecciona un foco —agua, suelo, biodiversidad, resiliencia climática o emisiones— y responde:")
    for item in [
        "¿Qué cambio muestra el escenario?",
        "¿Qué consecuencias podría tener para el territorio y para la vida de las juventudes?",
        "¿Qué municipios, ecosistemas o actores merecen especial atención?",
        "¿Qué información adicional sería necesaria antes de tomar decisiones?",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "Momento 6. Laboratorio de acción | 30 minutos", 2)
    add_body(document, "Propósito: convertir un hallazgo en una intervención juvenil concreta y verificable.", bold_prefix="Propósito:")
    add_body(document, "Cada grupo diseña una propuesta utilizando la siguiente fórmula:")
    add_callout(document, "FÓRMULA DE ACCIÓN", "Ante [problema o riesgo], jóvenes de [territorio] realizarán [acción concreta] con [aliados], comenzando por [primer paso] y verificando el avance mediante [indicador sencillo].", fill=GREEN_LIGHT, accent=GREEN)
    for item in [
        "La acción debe poder iniciar en un plazo máximo de tres meses.",
        "Debe reconocer al menos un actor comunitario o institucional aliado.",
        "Debe incluir una evidencia observable: personas vinculadas, nacimientos de agua monitoreados, contenidos producidos, parcelas demostrativas, acuerdos logrados u otra medida pertinente.",
        "Debe evitar promesas que dependan totalmente de recursos o decisiones externas.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "Momento 7. Galería, compromisos y cierre | 25 minutos", 2)
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Exponer las propuestas en formato galería; cada vocería dispone de dos minutos.",
        "Entregar a cada participante dos marcas de retroalimentación: una para la propuesta más viable y otra para la de mayor impacto territorial.",
        "Solicitar a cada grupo que defina un próximo paso, una persona enlace y una fecha tentativa de seguimiento.",
        "Cerrar con la pregunta: ¿qué idea me llevo y qué acción sí puedo empezar desde mi lugar?",
        "Aplicar la evaluación rápida de salida.",
    ]:
        add_numbered(document, item, decimal_id)

    add_heading(document, "10. Roles y responsabilidades", 1)
    add_table(
        document,
        ["Rol", "Antes de la actividad", "Durante y después"],
        [
            ("Coordinación general", "Define alcance, convocatoria, lugar y aliados; valida mensajes institucionales.", "Abre y cierra la jornada; acuerda seguimiento y custodia los productos."),
            ("Facilitación principal", "Adapta la metodología y prepara preguntas.", "Conduce tiempos, conversaciones y transiciones; garantiza participación equilibrada."),
            ("Apoyo técnico", "Prueba computador, proyector, conectividad y respaldo local del simulador.", "Opera o apoya el simulador y resuelve incidentes sin interrumpir la metodología."),
            ("Relatoría y logística", "Prepara materiales, registro y formatos de consentimiento.", "Documenta hallazgos, recoge fichas, aplica evaluación y organiza evidencias."),
            ("Enlace territorial", "Aporta contexto local y conecta con participantes y actores.", "Ayuda a interpretar referencias municipales y sostiene el seguimiento posterior."),
        ],
        [1800, 3500, 4060],
        header_fill=GREEN_LIGHT,
    )

    add_heading(document, "11. Materiales y condiciones", 1)
    materials = [
        "Computador con el simulador previamente cargado y cargador.",
        "Proyector o pantalla, extensión eléctrica y adaptadores necesarios.",
        "Conexión a internet; respaldo local preparado para contingencia.",
        "Mapa impreso o dibujado del sur del Huila.",
        "Notas adhesivas de dos colores, marcadores, cinta y papel tamaño pliego.",
        "Una ficha de trabajo por grupo y formatos de evaluación de salida.",
        "Agua y condiciones básicas de bienestar; espacio para pausa activa.",
        "Autorizaciones para fotografías o testimonios cuando corresponda.",
    ]
    for item in materials:
        add_bullet(document, item, bullet_id)

    add_heading(document, "12. Inclusión y cuidado", 1)
    for item in [
        "Usar lenguaje claro y explicar todo término técnico la primera vez que aparezca.",
        "Combinar intervenciones orales, escritura y trabajo visual para incluir distintas formas de participación.",
        "Garantizar accesibilidad física y pausas; consultar previamente necesidades específicas.",
        "Distribuir las vocerías y evitar que una sola persona maneje siempre la herramienta.",
        "No solicitar datos personales innecesarios ni publicar imágenes sin consentimiento.",
        "Evitar mensajes catastróficos; equilibrar riesgos, capacidades locales y posibilidades de acción.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "13. Riesgos operativos y contingencias", 1)
    add_table(
        document,
        ["Riesgo", "Prevención", "Respuesta durante la sesión", "Responsable"],
        [
            ("Falla de internet", "Abrir y probar el respaldo local antes de iniciar.", "Continuar con la cartografía e indicadores esenciales sin teselas externas.", "Apoyo técnico"),
            ("Falla de energía o proyección", "Llevar extensiones, adaptadores y capturas impresas de tres momentos.", "Trabajar con las láminas 2026, 2030 y 2035 y mantener la misma guía de preguntas.", "Coordinación / apoyo técnico"),
            ("Discusión dominada por pocas voces", "Definir roles y tiempos de intervención.", "Usar ronda breve, escritura individual y vocerías rotativas.", "Facilitación"),
            ("Interpretación del escenario como predicción exacta", "Repetir el mensaje metodológico en la introducción.", "Diferenciar en un tablero: estimación, observación y pregunta por verificar.", "Facilitación"),
            ("Propuestas demasiado amplias", "Presentar criterios de plazo, aliado y evidencia.", "Reducir la idea a un primer paso realizable en tres meses.", "Facilitación / enlace"),
        ],
        [1500, 2800, 3300, 1760],
        header_fill=GOLD_LIGHT,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 4,
    )

    add_heading(document, "14. Evaluación y seguimiento", 1)
    add_heading(document, "14.1 Evaluación de salida", 2)
    add_body(document, "Aplicar cuatro afirmaciones en escala de 1 a 5, donde 1 significa “nada” y 5 significa “mucho”:")
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Comprendo mejor la relación entre cambio climático y transformaciones del territorio.",
        "Puedo interpretar la línea temporal y los indicadores del simulador.",
        "La actividad incorporó experiencias y conocimientos de las juventudes.",
        "Identifiqué una acción concreta en la que podría participar.",
    ]:
        add_numbered(document, item, decimal_id)
    add_body(document, "Añadir dos preguntas abiertas: ¿qué fue lo más útil? y ¿qué debería cambiarse para una próxima sesión?")

    add_heading(document, "14.2 Indicadores de implementación", 2)
    add_table(
        document,
        ["Indicador", "Meta sugerida", "Medio de verificación"],
        [
            ("Participación", "Al menos 80 % de las personas permanece hasta el cierre.", "Listado de asistencia y registro de cierre."),
            ("Comprensión", "Al menos 75 % marca 4 o 5 en la primera afirmación.", "Evaluación de salida."),
            ("Producción", "Cada grupo formula una propuesta con primer paso e indicador.", "Fichas o carteles de acción."),
            ("Diversidad de voces", "Todos los grupos cuentan con vocería y relatoría diferenciadas.", "Observación de facilitación."),
            ("Seguimiento", "Se acuerda al menos una fecha o mecanismo de contacto posterior.", "Acta breve de compromisos."),
        ],
        [2500, 3300, 3560],
        header_fill=GREEN_LIGHT,
    )

    add_heading(document, "14.3 Seguimiento sugerido", 2)
    decimal_id = create_numbering_instance(document, decimal_abstract_id)
    for item in [
        "Dentro de los cinco días siguientes: consolidar fotografías autorizadas, fichas, evaluación y propuestas.",
        "Dentro de las dos semanas siguientes: devolver al grupo un resumen breve de hallazgos y próximos pasos.",
        "Dentro de los treinta días siguientes: realizar una reunión corta con enlaces juveniles y aliados para priorizar acciones.",
        "A los tres meses: revisar avances usando el indicador definido por cada grupo.",
    ]:
        add_numbered(document, item, decimal_id)

    document.add_page_break()

    add_heading(document, "Anexo A. Ficha de exploración y acción", 1)
    add_body(document, "Nombre del grupo: ________________________________    Municipio(s): ________________________________")
    add_prompt_box(document, 1, "Lectura inicial - 2026", "¿Qué caracteriza la huella cafetera y el estado de los indicadores al inicio del escenario?", 2)
    add_prompt_box(document, 2, "Cambio intermedio - 2030", "¿Qué expansión, retiro o variación de indicador llama más la atención?", 2)
    add_prompt_box(document, 3, "Escenario final - 2035", "¿Cuál es el principal riesgo u oportunidad que identifica el grupo?", 2)
    add_prompt_box(document, 4, "Contraste territorial", "¿Cómo se relaciona lo observado con experiencias o conocimientos del municipio? ¿Qué habría que verificar?", 3)
    add_prompt_box(document, 5, "Propuesta juvenil", "Complete: Ante [problema], jóvenes de [territorio] realizarán [acción] con [aliados], comenzando por [primer paso] y midiendo [indicador].", 4)

    annex_b_heading = add_heading(document, "Anexo B. Lista de verificación del equipo", 1)
    annex_b_heading.paragraph_format.page_break_before = True
    checks = [
        "La convocatoria informa propósito, duración, lugar y condiciones de participación.",
        "El espacio es accesible, ventilado y permite trabajo por grupos.",
        "El simulador fue probado en el equipo y navegador que se utilizarán.",
        "Existe un respaldo para conectividad, energía y proyección.",
        "Los materiales están separados por momento metodológico.",
        "Cada integrante del equipo conoce su rol y las señales de tiempo.",
        "El mensaje sobre el carácter estimativo del simulador está incorporado al guion.",
        "Se definió cómo se recogerán, custodiarán y devolverán los resultados.",
        "Las fotografías y testimonios cuentan con autorización.",
        "Existe una fecha o mecanismo de seguimiento posterior.",
    ]
    for item in checks:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(8)
        box = p.add_run("☐  ")
        set_run_font(box, name="Arial", size=11, color=GREEN, bold=True)
        text = p.add_run(item)
        set_run_font(text)

    add_heading(document, "Fuentes de referencia del simulador", 1)
    add_body(document, "La herramienta utiliza cartografía libre u oficial y resultados preprocesados. Para la presentación al público se recomienda mencionar las siguientes fuentes generales:")
    sources = [
        ("IGAC", "límites departamentales y municipales"),
        ("UPRA", "evaluaciones agropecuarias municipales y zonificación de aptitud para café"),
        ("IDEAM", "escenarios departamentales de cambio climático y redes de drenaje"),
        ("MapBiomas Colombia", "coberturas y transformaciones del suelo"),
        ("RUNAP - Parques Nacionales Naturales", "áreas protegidas"),
        ("OpenStreetMap / OpenFreeMap", "base cartográfica y atribución"),
        ("IPCC", "referentes para estimaciones de carbono por cambio de cobertura"),
    ]
    for source, use in sources:
        add_bullet(document, f"{source}: {use}.", bullet_id, bold_prefix=source + ":")
    add_callout(document, "NOTA FINAL", "Esta guía es una base de trabajo. El equipo podrá ajustar duración, número de grupos, lenguaje y énfasis municipal sin modificar el propósito central ni el mensaje sobre el carácter estimativo del simulador.", fill=GOLD_LIGHT, accent="8A651B")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
